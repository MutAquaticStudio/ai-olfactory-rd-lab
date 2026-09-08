import json

import numpy as np
import pandas as pd
import pytest

from olfactory.data_foundation import DataFoundationService
from olfactory.data_foundation.evidence import file_hash
from olfactory.data_foundation.public_panels import import_public_panel
from olfactory.training.dataset import load_versioned_snapshot
from olfactory.training.reliability import krippendorff_alpha_nominal, icc_2k
from olfactory.training.calibration import CalibrationBundle


def fixture_archive(tmp_path, approved=True):
    root = tmp_path / "archive"
    root.mkdir()
    pd.DataFrame([{"CID": 1, "CanonicalSMILES": "CCO"}]).to_csv(root / "molecules.csv", index=False)
    pd.DataFrame([{"Stimulus": 1, "CIDs": 1, "Concentration": 0.001, "Ratio": "1/1000", "Solvent": "paraffin oil"}]).to_csv(root / "stimuli.csv", index=False)
    pd.DataFrame([
        {"Stimulus": 1, "Subject": 1, "MeasurementValue": "SWEET", "Value": None},
        {"Stimulus": 1, "Subject": 2, "MeasurementValue": "SWEET", "Value": "0"},
        {"Stimulus": 1, "Subject": 2, "MeasurementValue": "MUSKY", "Value": "80"},
        {"Stimulus": 1, "Subject": 2, "MeasurementValue": "SWEET", "Value": "50"},
    ]).to_csv(root / "behavior.csv", index=False)
    files = ["molecules.csv", "stimuli.csv", "behavior.csv"]
    (root / "source-manifest.json").write_text(json.dumps({"archive": "keller_2016", "pyrfume_data_commit": "a" * 40,
        "files": [{"filename": name, "sha256": file_hash(root / name)} for name in files]}))
    registry = tmp_path / "registry.json"
    registry.write_text(json.dumps({"schema_version": 1, "pyrfume_data_commit": "a" * 40,
        "archives": {"keller_2016": {"license_status": "APPROVED" if approved else "REVIEW_REQUIRED", "files": files}}}))
    return root, registry


def test_panel_import_keeps_missing_and_replicates_unknown_and_musky_in_review(tmp_path):
    root, registry = fixture_archive(tmp_path)
    service = DataFoundationService(tmp_path / "data", ("sweet", "musk"))
    result = import_public_panel(service, root, registry_path=registry, dataset_version="panel-fixture", chunk_size=2)
    path = service.snapshot_root / "panel-fixture"
    frame = pd.read_parquet(path / "ratings.parquet")
    assert len(frame) == 4
    assert frame.presence_state.eq("UNASSESSED").all()
    assert frame.replicate_number.isna().all()
    assert frame.intensity.isna().all()
    assert frame.loc[frame.source_term == "MUSKY", "descriptor"].isna().all()
    assert result["report"]["counts"]["explicit_zero_ratings"] == 1
    assert frame.source_row.nunique() == 4  # repeated source observations are preserved
    assert len(service.list_snapshots()) == 1
    with pytest.raises(ValueError, match="AUDIT_ONLY"):
        load_versioned_snapshot(path / "ratings.parquet", ("sweet", "musk"))


def test_unapproved_panel_never_reads_data_files(tmp_path):
    root, registry = fixture_archive(tmp_path, approved=False)
    (root / "behavior.csv").unlink()
    service = DataFoundationService(tmp_path / "data", ("sweet",))
    with pytest.raises(PermissionError, match="LICENSE_REVIEW_REQUIRED"):
        import_public_panel(service, root, registry_path=registry, dataset_version="blocked")
    assert service.list_snapshots() == []


def test_panel_checksum_mismatch_cannot_create_release(tmp_path):
    root, registry = fixture_archive(tmp_path)
    (root / "behavior.csv").write_text("corrupt")
    service = DataFoundationService(tmp_path / "data", ("sweet",))
    with pytest.raises(ValueError, match="checksum"):
        import_public_panel(service, root, registry_path=registry, dataset_version="corrupt")
    assert service.list_snapshots() == []


def test_reliability_matches_hand_calculated_coincidences_and_absolute_agreement():
    # Seven pairable ratings: Do=2/7 and De=4/7. Singleton is excluded.
    ratings = np.array([[1, 1, 1], [0, 0, np.nan], [1, 0, np.nan], [1, np.nan, np.nan]])
    assert krippendorff_alpha_nominal(ratings) == pytest.approx(0.5)
    assert np.isnan(krippendorff_alpha_nominal(np.ones((3, 3))))
    # MS_rows=8, MS_columns=1.5, MS_error=0, n=3.
    assert icc_2k(np.array([[1, 2], [3, 4], [5, 6]])) == pytest.approx(16 / 17)
    assert np.isnan(icc_2k(np.ones((1, 3))))


def test_calibration_cannot_label_one_class_or_unknown_as_calibrated():
    logits = np.zeros((80, 2))
    targets = np.column_stack([np.ones(80), np.full(80, np.nan)])
    bundle = CalibrationBundle.fit(logits, targets, ("sweet", "musk"), mask=np.ones_like(targets, dtype=bool))
    assert all(method == "UNCALIBRATED_INSUFFICIENT_EVIDENCE" for method in bundle.methods)


def test_aggregate_adapter_keeps_use_and_applicability_separate(tmp_path):
    root, registry_path = fixture_archive(tmp_path)
    pd.DataFrame([{"Stimulus": "sample-low", "CID": "1", "Conc": "low"}]).to_csv(root / "stimuli.csv", index=False)
    for filename, value in (("behavior_1.csv", 32), ("behavior_2.csv", 0)):
        pd.DataFrame([{"Stimulus": "sample-low", "Sweet": value}]).to_csv(root / filename, index=False)
    files = ["molecules.csv", "stimuli.csv", "behavior_1.csv", "behavior_2.csv"]
    registry_path.write_text(json.dumps({"schema_version": 1, "pyrfume_data_commit": "a" * 40,
        "archives": {"dravnieks_1985": {"license_status": "APPROVED", "files": files}}}))
    (root / "source-manifest.json").write_text(json.dumps({"archive": "dravnieks_1985", "pyrfume_data_commit": "a" * 40,
        "files": [{"filename": name, "sha256": file_hash(root / name)} for name in files]}))
    service = DataFoundationService(tmp_path / "data", ("sweet",))
    import_public_panel(service, root, registry_path=registry_path, dataset_version="aggregate-fixture")
    frame = pd.read_parquet(service.snapshot_root / "aggregate-fixture/ratings.parquet")
    assert len(frame) == 2
    assert frame.measurement_table.nunique() == 2
    assert frame.assessor_id.isna().all()
    assert frame.presence_state.eq("UNASSESSED").all()
    assert frame.inchikey.nunique() == 1


def test_training_loader_refuses_to_pool_measurement_conditions(tmp_path):
    frame = pd.DataFrame([{
        "assessment_id": str(i), "study_name": "fixture", "session_name": "session-1",
        "assessor_id": "assessor", "inchikey": "fixture-key", "descriptor": "sweet",
        "presence_state": "PRESENT", "intensity": 5., "replicate_number": 1,
        "stereo_state": "ACHIRAL", "isomeric_smiles": "CCO", "concentration": concentration,
    } for i, concentration in enumerate((0.001, 0.00001))])
    path = tmp_path / "different-conditions.parquet"
    frame.to_parquet(path)
    with pytest.raises(ValueError, match="measurement_domain"):
        load_versioned_snapshot(path, ("sweet",))


def fixture_workbook(tmp_path):
    from openpyxl import Workbook
    root, registry_path = fixture_archive(tmp_path)
    terms = ["CAN OR CAN'T SMELL", "KNOW OR DON'T KNOW THE SMELL", "THE ODOR IS:",
             "HOW STRONG IS THE SMELL?", "HOW PLEASANT IS THE SMELL?", "HOW FAMILIAR IS THE SMELL?",
             "EDIBLE", "BAKERY", "SWEET", "FRUIT", "FISH", "GARLIC", "SPICES", "COLD", "SOUR",
             "BURNT", "ACID", "WARM", "MUSKY", "SWEATY", "AMMONIA/URINOUS", "DECAYED", "WOOD",
             "GRASS", "FLOWER", "CHEMICAL"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "data"
    sheet.append(["Fixture, not real panel data"])
    sheet.append(["Source notes"])
    sheet.append(["C.A.S.", "Catalogue #*", "CID", "Odor", "Odor dilution",
                  "Subject # (this study)", "VIAL #", *terms])
    rows = []
    for i, name in enumerate(("ethanol", "ethanol (replicate)")):
        values = ["I smell something" if i == 0 else "I can't smell anything", None,
                  "1" if i == 0 else None, 0 if i == 0 else None, None, None, *([None] * 20)]
        sheet.append(["64-17-5", "fixture-catalog", 1, name, "1/1,000", 1, 10 + i, *values])
        rows.extend({"Stimulus": "1", "Subject": "1", "MeasurementValue": term, "Value": val}
                    for term, val in zip(terms, values))
    path = root / "original.xlsx"
    workbook.save(path)
    pd.DataFrame(rows).to_csv(root / "behavior.csv", index=False)
    manifest = json.loads((root / "source-manifest.json").read_text())
    for item in manifest["files"]:
        item["sha256"] = file_hash(root / item["filename"])
    (root / "source-manifest.json").write_text(json.dumps(manifest))
    registry = json.loads(registry_path.read_text())
    registry["archives"]["keller_2016"]["original_workbook"] = {
        "filename": path.name, "sha256": file_hash(path), "url": "https://example.invalid/fixture.xlsx"}
    registry_path.write_text(json.dumps(registry))
    return root, registry_path, path


def test_original_trials_recover_vials_without_approving_ratings(tmp_path):
    root, registry, workbook = fixture_workbook(tmp_path)
    service = DataFoundationService(tmp_path / "data", ("sweet", "musk"))
    result = import_public_panel(service, root, registry_path=registry, dataset_version="recovered",
                                 original_workbook=workbook.name, chunk_size=7)
    path = service.snapshot_root / "recovered"
    frame = pd.read_parquet(path / "ratings.parquet")
    assert len(frame) == 52 and frame.trial_id.nunique() == 2
    assert set(frame.vial_id) == {"10", "11"}
    assert set(frame.source_excel_row) == {"4", "5"}
    assert frame.loc[frame.vial_id == "11", "source_replicate_marker"].eq("true").all()
    assert frame.session_id.isna().all() and frame.replicate_number.isna().all()
    assert frame.presence_state.eq("UNASSESSED").all() and frame.intensity.isna().all()
    assert frame.loc[frame.source_term == "MUSKY", "descriptor"].isna().all()
    sweet = frame[frame.source_term == "SWEET"]
    assert list(sweet.blank_reason) == ["UNRECORDED_DESCRIPTOR_VALUE", "SKIPPED_NO_DETECTION"]
    assert frame.loc[frame.source_term == "THE ODOR IS:", "raw_numeric_value"].isna().all()
    assert result["report"]["counts"]["explicit_zero_descriptor_ratings"] == 0
    assert result["report"]["counts"]["explicit_zero_other_ratings"] == 1
    assert result["report"]["counts"]["recovered_trials"] == 2
    manifest = json.loads((path / "manifest.json").read_text())
    assert manifest["original_workbook"]["sha256"] == file_hash(workbook)
    assert not manifest["training_eligible"]
    import_public_panel(service, root, registry_path=registry, dataset_version="same-source-new-version",
                        original_workbook=workbook.name, chunk_size=13)
    other = pd.read_parquet(service.snapshot_root / "same-source-new-version/ratings.parquet")
    assert list(frame.trial_id) == list(other.trial_id)
    assert list(frame.source_excel_row) == list(other.source_excel_row)
    with pytest.raises(FileExistsError):
        import_public_panel(service, root, registry_path=registry, dataset_version="recovered",
                            original_workbook=workbook.name)


@pytest.mark.parametrize("mutation", ["value", "subject", "stimulus", "truncated", "extra"])
def test_original_mismatch_never_publishes_snapshot(tmp_path, mutation):
    root, registry, workbook = fixture_workbook(tmp_path)
    frame = pd.read_csv(root / "behavior.csv", dtype=str)
    if mutation == "value":
        frame.loc[0, "Value"] = "different"
    elif mutation == "subject":
        frame.loc[0, "Subject"] = "2"
    elif mutation == "stimulus":
        frame.loc[0, "Stimulus"] = "999"
    elif mutation == "truncated":
        frame = frame.iloc[:-1]
    else:
        frame = pd.concat([frame, frame.iloc[:1]])
    frame.to_csv(root / "behavior.csv", index=False)
    manifest = json.loads((root / "source-manifest.json").read_text())
    for item in manifest["files"]:
        item["sha256"] = file_hash(root / item["filename"])
    (root / "source-manifest.json").write_text(json.dumps(manifest))
    service = DataFoundationService(tmp_path / "data", ("sweet", "musk"))
    with pytest.raises(ValueError, match="Keller.*mismatch"):
        import_public_panel(service, root, registry_path=registry, dataset_version="mismatch",
                            original_workbook=workbook.name, chunk_size=7)
    assert not service.list_snapshots()
    assert not list(service.snapshot_root.iterdir())


def test_original_checksum_required(tmp_path):
    root, registry, workbook = fixture_workbook(tmp_path)
    workbook.write_bytes(b"changed")
    service = DataFoundationService(tmp_path / "data", ("sweet",))
    with pytest.raises(ValueError, match="checksum"):
        import_public_panel(service, root, registry_path=registry, dataset_version="bad",
                            original_workbook=workbook.name)
    assert not service.list_snapshots()


def repin_workbook(registry_path, workbook):
    registry = json.loads(registry_path.read_text())
    registry["archives"]["keller_2016"]["original_workbook"]["sha256"] = file_hash(workbook)
    registry_path.write_text(json.dumps(registry))


def test_original_duplicate_subject_vial_is_rejected(tmp_path):
    from openpyxl import load_workbook
    root, registry, path = fixture_workbook(tmp_path)
    workbook = load_workbook(path)
    workbook["data"].cell(5, 7).value = 10
    workbook.save(path)
    repin_workbook(registry, path)
    service = DataFoundationService(tmp_path / "data", ("sweet",))
    with pytest.raises(ValueError, match="duplicate trial identity"):
        import_public_panel(service, root, registry_path=registry, dataset_version="duplicate",
                            original_workbook=path.name)
    assert not service.list_snapshots()


def test_original_cas_in_cid_field_is_preserved_not_repaired(tmp_path):
    from openpyxl import load_workbook
    root, registry, path = fixture_workbook(tmp_path)
    workbook = load_workbook(path)
    for row in (4, 5):
        workbook["data"].cell(row, 3).value = "64-17-5"
    workbook.save(path)
    repin_workbook(registry, path)
    stimuli = pd.read_csv(root / "stimuli.csv", dtype=str)
    stimuli["CIDs"] = "64-17-5"
    stimuli.to_csv(root / "stimuli.csv", index=False)
    manifest = json.loads((root / "source-manifest.json").read_text())
    for item in manifest["files"]:
        item["sha256"] = file_hash(root / item["filename"])
    (root / "source-manifest.json").write_text(json.dumps(manifest))
    service = DataFoundationService(tmp_path / "data", ("sweet", "musk"))
    result = import_public_panel(service, root, registry_path=registry, dataset_version="cas",
                                 original_workbook=path.name)
    frame = pd.read_parquet(service.snapshot_root / "cas/ratings.parquet")
    assert frame.raw_source_identifier.eq("64-17-5").all()
    assert frame.source_identifier.eq("64-17-5").all()
    assert frame.identifier_type.eq("CAS_IN_SOURCE_CID_FIELD").all()
    assert frame.isomeric_smiles.isna().all()
    assert "SOURCE_CID_IS_CAS_REVIEW_REQUIRED" in result["report"]["warnings"]
    assert result["report"]["warnings"]["STRUCTURE_REVIEW_REQUIRED"] == 52
    assert result["report"]["warnings"]["MUSKY_SEMANTIC_CONFLICT_WITH_PERFUMERY_MUSK"] == 2
    queue = pd.read_csv(service.snapshot_root / "cas/structure_review_queue.csv")
    assert len(queue) == 1 and queue.measurement_rows.iloc[0] == 52


def test_original_path_cannot_escape_archive(tmp_path):
    root, registry, path = fixture_workbook(tmp_path)
    outside = tmp_path / "outside.xlsx"
    outside.write_bytes(path.read_bytes())
    path.unlink()
    path.symlink_to(outside)
    service = DataFoundationService(tmp_path / "data", ("sweet",))
    with pytest.raises(ValueError, match="outside"):
        import_public_panel(service, root, registry_path=registry, dataset_version="outside",
                            original_workbook=path.name)
    assert not service.list_snapshots()


def test_cas_correction_is_versioned_and_does_not_approve_evidence(tmp_path):
    root, registry, path = fixture_workbook(tmp_path)
    service = DataFoundationService(tmp_path / "data", ("sweet", "musk"))
    import_public_panel(service, root, registry_path=registry, dataset_version="before-cas",
                        original_workbook=path.name)
    parent = service.snapshot_root / "before-cas/manifest.json"
    before_hash = file_hash(parent)
    correction = {
        "correction_id": "fixture-cas-correction", "parent_dataset_version": "before-cas",
        "parent_manifest_sha256": before_hash,
        "original_workbook_sha256": file_hash(path),
        "old_cas": "64-17-5", "new_cas": "110-19-0", "source_identifier": "1",
        "source_names": ["ethanol", "ethanol (replicate)"],
        "expected_measurement_rows": 52, "expected_trials": 2,
        "authorization": "EXPLICIT_USER_CAS_CORRECTION", "authorized_at": "2026-09-08",
        "reason": "Synthetic correction fixture only", "evidence_urls": ["https://example.invalid/fixture"],
    }
    result = import_public_panel(service, root, registry_path=registry, dataset_version="after-cas",
                                 original_workbook=path.name, cas_correction=correction, chunk_size=7)
    old = pd.read_parquet(service.snapshot_root / "before-cas/ratings.parquet")
    new = pd.read_parquet(service.snapshot_root / "after-cas/ratings.parquet")
    assert new.source_cas.eq("110-19-0").all()
    assert new.raw_source_cas.eq("64-17-5").all()
    assert new.cas_correction_id.eq("fixture-cas-correction").all()
    assert old.drop(columns="source_cas").equals(new[old.columns].drop(columns="source_cas"))
    assert file_hash(parent) == before_hash
    assert new.evidence_review.eq("UNREVIEWED").all()
    assert new.presence_state.eq("UNASSESSED").all()
    assert result["report"]["counts"]["cas_corrected_rows"] == 52
    child = json.loads((service.snapshot_root / "after-cas/manifest.json").read_text())
    assert child["cas_correction"]["parent_manifest_sha256"] == before_hash
    assert not child["training_eligible"]
    with pytest.raises(FileExistsError):
        import_public_panel(service, root, registry_path=registry, dataset_version="after-cas",
                            original_workbook=path.name, cas_correction=correction)
    for field, value in (("parent_manifest_sha256", "0" * 64), ("expected_measurement_rows", 51),
                         ("new_cas", "110-19-1"), ("authorization", "PENDING")):
        with pytest.raises((ValueError, PermissionError)):
            import_public_panel(service, root, registry_path=registry, dataset_version="rejected-" + field,
                                original_workbook=path.name, cas_correction={**correction, field: value})
        assert not (service.snapshot_root / ("rejected-" + field)).exists()
